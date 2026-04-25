from rest_framework import viewsets
from rest_framework.decorators import action
from rest_framework.permissions import AllowAny
from rest_framework.response import Response
from django.http import HttpResponse
from .models import Segmento, Campana, Lead
from .serializers import SegmentoSerializer, CampanaSerializer, LeadSerializer
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.drawing.image import Image as XLImage
import os
from django.conf import settings
from datetime import datetime

class SegmentoViewSet(viewsets.ModelViewSet):
    queryset = Segmento.objects.all()
    serializer_class = SegmentoSerializer
    permission_classes = [AllowAny]

class CampanaViewSet(viewsets.ModelViewSet):
    queryset = Campana.objects.all().order_by('-fecha_inicio')
    serializer_class = CampanaSerializer
    permission_classes = [AllowAny]

class LeadViewSet(viewsets.ModelViewSet):
    queryset = Lead.objects.all().order_by('-fecha_creacion')
    serializer_class = LeadSerializer
    permission_classes = [AllowAny]

    @action(detail=True, methods=['post'])
    def convertir_a_cliente(self, request, pk=None):
        lead = self.get_object()
        cliente = lead.convertir_a_cliente()
        serializer = self.get_serializer(lead)
        return Response({
            'message': 'Lead convertido a cliente exitosamente',
            'lead': serializer.data,
            'cliente_id': cliente.id
        })

    @action(detail=False, methods=['get'])
    def excel(self, request):
        leads = self.get_queryset()
        
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Leads"
        
        # Styles
        header_font = Font(name='Arial', size=12, bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
        bold_font = Font(bold=True)
        center_align = Alignment(horizontal='center', vertical='center')
        thin_border = Border(left=Side(style='thin'), right=Side(style='thin'), top=Side(style='thin'), bottom=Side(style='thin'))
        
        # Logo
        logo_path = os.path.join(settings.BASE_DIR, 'logo_kave.png')
        if os.path.exists(logo_path):
            img = XLImage(logo_path)
            img.width = 180
            img.height = 80
            ws.add_image(img, 'A1')
            
        # Company Info
        ws.merge_cells('C1:G1')
        ws['C1'] = "TRANSFORMADORES KAVE S.A.S."
        ws['C1'].font = Font(size=18, bold=True, color="1E3A8A")
        ws['C1'].alignment = Alignment(horizontal='center')
        
        ws.merge_cells('C2:G2')
        ws['C2'] = "Lista de Leads de Marketing"
        ws['C2'].alignment = Alignment(horizontal='center')
        
        ws['A5'] = f"Total de leads: {leads.count()}"
        ws.merge_cells('A5:G5')
        
        # Table Headers
        headers = ['Nombre', 'Email', 'Teléfono', 'Empresa', 'Fuente', 'Campaña', 'Estado', 'Puntuación', 'Fecha Creación']
        col_widths = [25, 30, 15, 25, 15, 25, 15, 12, 18]
        
        for i, header in enumerate(headers, 1):
            cell = ws.cell(row=7, column=i)
            cell.value = header
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = center_align
            cell.border = thin_border
            ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = col_widths[i-1]
        
        # Data rows
        for row_idx, lead in enumerate(leads, 8):
            ws.cell(row=row_idx, column=1).value = lead.nombre
            ws.cell(row=row_idx, column=2).value = lead.email
            ws.cell(row=row_idx, column=3).value = lead.telefono
            ws.cell(row=row_idx, column=4).value = lead.empresa
            ws.cell(row=row_idx, column=5).value = lead.get_fuente_display()
            ws.cell(row=row_idx, column=6).value = lead.campana.nombre if lead.campana else ''
            ws.cell(row=row_idx, column=7).value = lead.get_estado_display()
            ws.cell(row=row_idx, column=8).value = lead.puntuacion
            ws.cell(row=row_idx, column=9).value = lead.fecha_creacion.strftime("%Y-%m-%d %H:%M")
            
            # Apply borders
            for col in range(1, 10):
                ws.cell(row=row_idx, column=col).border = thin_border
        
        # Footer
        footer_row = len(leads) + 10
        ws.cell(row=footer_row, column=1).value = f"Generado el: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}"
        ws.merge_cells(f'A{footer_row}:G{footer_row}')
        
        response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
        response['Content-Disposition'] = 'attachment; filename=leads.xlsx'
        
        wb.save(response)
        return response
        leads = self.get_queryset()
        
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Leads"
        
        # Styles
        header_font = Font(name='Arial', size=12, bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
        bold_font = Font(bold=True)
        center_align = Alignment(horizontal='center', vertical='center')
        thin_border = Border(left=Side(style='thin'), right=Side(style='thin'), top=Side(style='thin'), bottom=Side(style='thin'))
        
        # Logo
        logo_path = os.path.join(settings.BASE_DIR, 'logo_kave.png')
        if os.path.exists(logo_path):
            img = XLImage(logo_path)
            img.width = 180
            img.height = 80
            ws.add_image(img, 'A1')
            
        # Company Info
        ws.merge_cells('C1:G1')
        ws['C1'] = "TRANSFORMADORES KAVE S.A.S."
        ws['C1'].font = Font(size=18, bold=True, color="1E3A8A")
        ws['C1'].alignment = Alignment(horizontal='center')
        
        ws.merge_cells('C2:G2')
        ws['C2'] = "Lista de Leads de Marketing"
        ws['C2'].alignment = Alignment(horizontal='center')
        
        ws['A5'] = f"Total de leads: {leads.count()}"
        ws.merge_cells('A5:G5')
        
        # Table Headers
        headers = ['Nombre', 'Email', 'Teléfono', 'Empresa', 'Fuente', 'Campaña', 'Estado', 'Puntuación', 'Fecha Creación']
        col_widths = [25, 30, 15, 25, 15, 25, 15, 12, 18]
        
        for i, header in enumerate(headers, 1):
            cell = ws.cell(row=7, column=i)
            cell.value = header
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = center_align
            cell.border = thin_border
            ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = col_widths[i-1]
        
        # Data rows
        for row_idx, lead in enumerate(leads, 8):
            ws.cell(row=row_idx, column=1).value = lead.nombre
            ws.cell(row=row_idx, column=2).value = lead.email
            ws.cell(row=row_idx, column=3).value = lead.telefono
            ws.cell(row=row_idx, column=4).value = lead.empresa
            ws.cell(row=row_idx, column=5).value = lead.get_fuente_display()
            ws.cell(row=row_idx, column=6).value = lead.campana.nombre if lead.campana else ''
            ws.cell(row=row_idx, column=7).value = lead.get_estado_display()
            ws.cell(row=row_idx, column=8).value = lead.puntuacion
            ws.cell(row=row_idx, column=9).value = lead.fecha_creacion.strftime("%Y-%m-%d %H:%M")
            
            # Apply borders
            for col in range(1, 10):
                ws.cell(row=row_idx, column=col).border = thin_border
        
        # Footer
        footer_row = len(leads) + 10
        ws.cell(row=footer_row, column=1).value = f"Generado el: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}"
        ws.merge_cells(f'A{footer_row}:G{footer_row}')
        
        response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
        response['Content-Disposition'] = 'attachment; filename=leads.xlsx'
        
        wb.save(response)
        return response
