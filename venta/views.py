from rest_framework import viewsets, permissions, status
from rest_framework.decorators import action
from rest_framework.response import Response
from django.http import HttpResponse
from .models import OrdenVenta, DetalleOrdenVenta, FacturaVenta, NotaCredito, DetalleNotaCredito
from .serializers import OrdenVentaSerializer, DetalleOrdenVentaSerializer, FacturaVentaSerializer, NotaCreditoSerializer, DetalleNotaCreditoSerializer

class OrdenVentaViewSet(viewsets.ModelViewSet):
    queryset = OrdenVenta.objects.all()
    serializer_class = OrdenVentaSerializer

class DetalleOrdenVentaViewSet(viewsets.ModelViewSet):
    queryset = DetalleOrdenVenta.objects.all()
    serializer_class = DetalleOrdenVentaSerializer

class FacturaVentaViewSet(viewsets.ModelViewSet):
    queryset = FacturaVenta.objects.all()
    serializer_class = FacturaVentaSerializer

    @action(detail=True, methods=['get'])
    def export_pdf(self, request, pk=None):
        from io import BytesIO
        from reportlab.pdfgen import canvas
        from reportlab.lib.pagesizes import letter
        from reportlab.lib import colors
        from django.http import HttpResponse

        factura = self.get_object()
        cliente = factura.orden_venta.cliente
        detalles = factura.orden_venta.detalles.all()

        buffer = BytesIO()
        p = canvas.Canvas(buffer, pagesize=letter)
        width, height = letter

        # Encabezado
        p.setFillColor(colors.HexColor("#1e293b"))
        p.rect(0, height - 80, width, 80, fill=1)
        p.setFillColor(colors.white)
        p.setFont("Helvetica-Bold", 16)
        p.drawString(40, height - 45, "8AMPERIOS ERP - FACTURA DE VENTA")
        p.setFont("Helvetica", 10)
        p.drawString(40, height - 60, f"Factura N°: {factura.numero_factura}")

        # Info Cliente
        p.setFillColor(colors.black)
        p.setFont("Helvetica-Bold", 12)
        p.drawString(40, height - 120, f"Cliente: {cliente.nombre}")
        p.setFont("Helvetica", 10)
        p.drawString(40, height - 135, f"Email: {cliente.email}")
        p.drawString(40, height - 150, f"Fecha Emisión: {factura.fecha_emision}")

        # Tabla Detalles
        p.line(40, height - 170, width - 40, height - 170)
        p.setFont("Helvetica-Bold", 10)
        p.drawString(50, height - 185, "PRODUCTO")
        p.drawRightString(width - 250, height - 185, "CANT")
        p.drawRightString(width - 150, height - 185, "UNITARIO")
        p.drawRightString(width - 50, height - 185, "TOTAL")
        p.line(40, height - 195, width - 40, height - 195)

        y = height - 215
        p.setFont("Helvetica", 10)
        for d in detalles:
            p.drawString(50, y, d.producto.nombre[:40])
            p.drawRightString(width - 250, y, str(d.cantidad))
            p.drawRightString(width - 150, y, f"${d.precio_unitario:,.2f}")
            p.drawRightString(width - 50, y, f"${(d.cantidad * d.precio_unitario):,.2f}")
            y -= 20
            if y < 50:
                p.showPage()
                y = height - 50

        p.line(40, y, width - 40, y)
        y -= 25
        p.setFont("Helvetica-Bold", 12)
        p.drawString(width - 200, y, "TOTAL:")
        p.drawRightString(width - 50, y, f"${factura.total:,.2f}")

        p.showPage()
        p.save()
        buffer.seek(0)
        response = HttpResponse(buffer, content_type='application/pdf')
        response['Content-Disposition'] = f'attachment; filename="Factura_{factura.numero_factura}.pdf"'
        return response

    @action(detail=True, methods=['get'])
    def export_excel(self, request, pk=None):
        import openpyxl
        from django.http import HttpResponse
        from io import BytesIO

        factura = self.get_object()
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Detalle Factura"

        # Encabezado
        ws['A1'] = "ERP 8AMPERIOS - FACTURA DE VENTA"
        ws['A2'] = f"Factura N°: {factura.numero_factura}"
        ws['A3'] = f"Cliente: {factura.orden_venta.cliente.nombre}"
        ws['A4'] = f"Fecha: {factura.fecha_emision}"

        # Columnas
        headers = ["Producto", "Cantidad", "Precio Unitario", "Subtotal"]
        ws.append([])
        ws.append(headers)

        for d in factura.orden_venta.detalles.all():
            ws.append([d.producto.nombre, d.cantidad, d.precio_unitario, d.cantidad * d.precio_unitario])

        ws.append([])
        ws.append(["", "", "TOTAL", factura.total])

        buffer = BytesIO()
        wb.save(buffer)
        buffer.seek(0)
        response = HttpResponse(buffer, content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
        response['Content-Disposition'] = f'attachment; filename="Factura_{factura.numero_factura}.xlsx"'
        return response


class NotaCreditoViewSet(viewsets.ModelViewSet):
    queryset = NotaCredito.objects.all()
    serializer_class = NotaCreditoSerializer
    
    @action(detail=True, methods=['get'])
    def export_pdf(self, request, pk=None):
        from io import BytesIO
        from reportlab.pdfgen import canvas
        from reportlab.lib.pagesizes import letter
        from reportlab.lib import colors
        from django.http import HttpResponse

        nc = self.get_object()
        factura = nc.factura
        cliente = factura.orden_venta.cliente
        detalles = nc.detalles.all()

        buffer = BytesIO()
        p = canvas.Canvas(buffer, pagesize=letter)
        width, height = letter

        # Encabezado
        p.setFillColor(colors.HexColor("#ef4444")) # Rojo para Notas de Crédito
        p.rect(0, height - 80, width, 80, fill=1)
        p.setFillColor(colors.white)
        p.setFont("Helvetica-Bold", 16)
        p.drawString(40, height - 45, "8AMPERIOS ERP - NOTA DE CRÉDITO")
        p.setFont("Helvetica", 10)
        p.drawString(40, height - 60, f"Nota N°: {nc.numero_nota} | Ref Factura: {factura.numero_factura}")

        # Info Cliente
        p.setFillColor(colors.black)
        p.setFont("Helvetica-Bold", 12)
        p.drawString(40, height - 120, f"Cliente: {cliente.nombre}")
        p.setFont("Helvetica", 10)
        p.drawString(40, height - 135, f"Tipo: {nc.get_tipo_display()}")
        p.drawString(40, height - 150, f"Fecha Emisión: {nc.fecha_emision}")
        p.drawString(40, height - 165, f"Motivo: {nc.motivo[:80]}...")

        # Tabla Detalles
        p.line(40, height - 185, width - 40, height - 185)
        p.setFont("Helvetica-Bold", 10)
        p.drawString(50, height - 200, "PRODUCTO")
        p.drawRightString(width - 250, height - 200, "CANT")
        p.drawRightString(width - 150, height - 200, "UNITARIO")
        p.drawRightString(width - 50, height - 200, "TOTAL")
        p.line(40, height - 210, width - 40, height - 210)

        y = height - 230
        p.setFont("Helvetica", 10)
        for d in detalles:
            p.drawString(50, y, d.producto.nombre[:40])
            p.drawRightString(width - 250, y, str(d.cantidad))
            p.drawRightString(width - 150, y, f"${d.precio_unitario:,.2f}")
            p.drawRightString(width - 50, y, f"${d.valor_total:,.2f}")
            y -= 20
            if y < 50:
                p.showPage()
                y = height - 50

        p.line(40, y, width - 40, y)
        y -= 25
        p.setFont("Helvetica-Bold", 12)
        p.drawString(width - 200, y, "VALOR IVA:")
        p.drawRightString(width - 50, y, f"${nc.valor_iva:,.2f}")
        y -= 20
        p.drawString(width - 200, y, "TOTAL NOTA:")
        p.drawRightString(width - 50, y, f"${nc.total:,.2f}")

        p.showPage()
        p.save()
        buffer.seek(0)
        response = HttpResponse(buffer, content_type='application/pdf')
        response['Content-Disposition'] = f'attachment; filename="NotaCredito_{nc.numero_nota}.pdf"'
        return response

    @action(detail=True, methods=['get'])
    def export_excel(self, request, pk=None):
        import openpyxl
        from django.http import HttpResponse
        from io import BytesIO

        nc = self.get_object()
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Detalle Nota Credito"

        ws['A1'] = "ERP 8AMPERIOS - NOTA DE CRÉDITO"
        ws['A2'] = f"Nota N°: {nc.numero_nota}"
        ws['A3'] = f"Factura Ref: {nc.factura.numero_factura}"
        ws['A4'] = f"Cliente: {nc.factura.orden_venta.cliente.nombre}"
        ws['A5'] = f"Motivo: {nc.motivo}"

        headers = ["Producto", "Cantidad", "Precio Unitario", "Subtotal"]
        ws.append([])
        ws.append(headers)

        for d in nc.detalles.all():
            ws.append([d.producto.nombre, d.cantidad, d.precio_unitario, d.valor_total])

        ws.append([])
        ws.append(["", "", "IVA", nc.valor_iva])
        ws.append(["", "", "TOTAL", nc.total])

        buffer = BytesIO()
        wb.save(buffer)
        buffer.seek(0)
        response = HttpResponse(buffer, content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
        response['Content-Disposition'] = f'attachment; filename="NotaCredito_{nc.numero_nota}.xlsx"'
        return response

    @action(detail=True, methods=['post'])
    def enviar_dian(self, request, pk=None):
        nota_credito = self.get_object()
        nota_credito.estado = 'enviada'
        nota_credito.save()
        return Response({'status': 'Nota de crédito enviada a DIAN'})
    
    @action(detail=True, methods=['post'])
    def aprobar(self, request, pk=None):
        nota_credito = self.get_object()
        nota_credito.estado = 'aprobada'
        nota_credito.save()
        return Response({'status': 'Nota de crédito aprobada'})


class DetalleNotaCreditoViewSet(viewsets.ModelViewSet):
    queryset = DetalleNotaCredito.objects.all()
    serializer_class = DetalleNotaCreditoSerializer