from rest_framework import viewsets, status, filters
from rest_framework.decorators import action
from rest_framework.response import Response
from django.http import HttpResponse
from .models import Cliente, Oportunidad, Cotizacion
from .serializers import ClienteSerializer, OportunidadSerializer, CotizacionSerializer
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.drawing.image import Image as XLImage
import os
from django.conf import settings
from datetime import datetime


class ClienteViewSet(viewsets.ModelViewSet):
    queryset = Cliente.objects.all().order_by('-fecha_registro')
    serializer_class = ClienteSerializer
    filter_backends = [filters.SearchFilter, filters.OrderingFilter]
    search_fields = ['nombre', 'nit', 'cedula', 'email', 'codigo_cliente', 'compania']
    ordering_fields = ['nombre', 'fecha_registro', 'estado', 'clasificacion']

    def get_queryset(self):
        qs = Cliente.objects.all().order_by('-fecha_registro')
        estado = self.request.query_params.get('estado')
        if estado:
            qs = qs.filter(estado=estado)
        return qs

    def destroy(self, request, *args, **kwargs):
        cliente = self.get_object()
        try:
            cliente.delete()
            return Response(status=status.HTTP_204_NO_CONTENT)
        except Exception as e:
            return Response(
                {
                    "error": f"No se puede eliminar el cliente '{cliente.nombre}' porque tiene registros asociados "
                             f"(cotizaciones, facturas u oportunidades). Elimine primero esos registros.",
                    "detalle": str(e)
                },
                status=status.HTTP_400_BAD_REQUEST
            )


class OportunidadViewSet(viewsets.ModelViewSet):
    serializer_class = OportunidadSerializer
    filter_backends = [filters.SearchFilter, filters.OrderingFilter]
    search_fields = ['titulo', 'descripcion']
    ordering_fields = ['titulo', 'estado', 'fecha_creacion', 'valor_estimado']

    def get_queryset(self):
        qs = Oportunidad.objects.all().order_by('-fecha_creacion')
        cliente_id = self.request.query_params.get('cliente')
        if cliente_id:
            qs = qs.filter(cliente_id=cliente_id)
        estado = self.request.query_params.get('estado')
        if estado:
            qs = qs.filter(estado=estado)
        return qs


class CotizacionViewSet(viewsets.ModelViewSet):
    serializer_class = CotizacionSerializer
    filter_backends = [filters.SearchFilter, filters.OrderingFilter]
    search_fields = ['numero_cotizacion', 'asunto']
    ordering_fields = ['numero_cotizacion', 'estado', 'fecha_creacion', 'gran_total']

    def get_queryset(self):
        qs = Cotizacion.objects.all().order_by('-fecha_creacion')
        cliente_id = self.request.query_params.get('cliente')
        if cliente_id:
            qs = qs.filter(cliente_id=cliente_id)
        estado = self.request.query_params.get('estado')
        if estado:
            qs = qs.filter(estado=estado)
        return qs

    @action(detail=True, methods=['get'])
    def excel(self, request, pk=None):
        cotizacion = self.get_object()
        
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Cotización"
        
        # Styles
        header_font = Font(name='Arial', size=12, bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
        bold_font = Font(bold=True)
        center_align = Alignment(horizontal='center', vertical='center')
        thin_border = Border(left=Side(style='thin'), right=Side(style='thin'), top=Side(style='thin'), bottom=Side(style='thin'))
        
        # Logo
        logo_path = os.path.join(settings.BASE_DIR, 'logo_kave.png')
        if os.path.exists(logo_path):
            img = XLImage(logo_path)
            img.width = 180
            img.height = 80
            ws.add_image(img, 'A1')
            
        # Company Info
        ws.merge_cells('C1:F1')
        ws['C1'] = "TRANSFORMADORES KAVE S.A.S."
        ws['C1'].font = Font(size=18, bold=True, color="1E3A8A")
        ws['C1'].alignment = Alignment(horizontal='center')
        
        ws.merge_cells('C2:F2')
        ws['C2'] = "Innovando en Transformadores"
        ws['C2'].alignment = Alignment(horizontal='center')
        
        ws['A5'] = "A continuación anexamos nuestra propuesta la cual esperamos sea de su completo agrado:"
        ws.merge_cells('A5:F5')
        
        # Cotizacion Info
        ws['A8'] = "Cotización N°:"
        ws['B8'] = cotizacion.numero_cotizacion or str(cotizacion.id)
        ws['A8'].font = bold_font
        
        ws['A9'] = "Fecha:"
        ws['B9'] = cotizacion.fecha_creacion.strftime("%Y-%m-%d")
        ws['A9'].font = bold_font
        
        ws['D8'] = "Cliente:"
        ws['E8'] = cotizacion.cliente.nombre
        ws['D8'].font = bold_font
        
        ws['D9'] = "NIT/Cédula:"
        ws['E9'] = cotizacion.cliente.cedula or "N/A"
        ws['D9'].font = bold_font

        ws['A11'] = "Asunto:"
        ws.merge_cells('B11:F11')
        ws['B11'] = cotizacion.asunto
        ws['A11'].font = bold_font
        
        # Table Headers
        headers = ['Item', 'Descripción / Producto', 'Unidad', 'Cantidad', 'V/Unitario', 'V/Total']
        col_widths = [8, 45, 12, 12, 18, 18]
        
        for i, header in enumerate(headers, 1):
            cell = ws.cell(row=13, column=i)
            cell.value = header
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = center_align
            cell.border = thin_border
            ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = col_widths[i-1]
            
        # Table Items
        row = 14
        for detalle in cotizacion.detalles.all().order_by('item'):
            ws.cell(row=row, column=1, value=detalle.item).alignment = center_align
            ws.cell(row=row, column=2, value=detalle.producto)
            ws.cell(row=row, column=3, value=detalle.unidad).alignment = center_align
            ws.cell(row=row, column=4, value=detalle.cantidad).alignment = center_align
            
            v_unit = ws.cell(row=row, column=5, value=float(detalle.valor_unitario))
            v_unit.number_format = '"$"#,##0'
            
            v_total = ws.cell(row=row, column=6, value=float(detalle.valor_total))
            v_total.number_format = '"$"#,##0'
            
            for i in range(1, 7):
                ws.cell(row=row, column=i).border = thin_border
            row += 1
            
        # Totals
        row += 1
        ws.cell(row=row, column=5, value="SUBTOTAL:").font = bold_font
        subt = ws.cell(row=row, column=6, value=float(cotizacion.valor_total))
        subt.font = bold_font
        subt.number_format = '"$"#,##0'
        
        row += 1
        ws.cell(row=row, column=5, value=f"% IVA ({cotizacion.porcentaje_iva}%):").font = bold_font
        iva_val = float(cotizacion.valor_total) * float(cotizacion.porcentaje_iva) / 100
        iva_cell = ws.cell(row=row, column=6, value=iva_val)
        iva_cell.font = bold_font
        iva_cell.number_format = '"$"#,##0'
        
        row += 1
        ws.cell(row=row, column=5, value="GRAN TOTAL:").font = bold_font
        ws.cell(row=row, column=5).fill = PatternFill(start_color="D1D5DB", end_color="D1D5DB", fill_type="solid")
        g_total = ws.cell(row=row, column=6, value=float(cotizacion.gran_total))
        g_total.font = bold_font
        g_total.fill = PatternFill(start_color="D1D5DB", end_color="D1D5DB", fill_type="solid")
        g_total.number_format = '"$"#,##0'

        # Commercial Conditions
        cond_row = row + 2
        ws.cell(row=cond_row, column=2, value="CONDICIONES COMERCIALES").font = Font(bold=True, color="1E3A8A")
        
        cond_row += 1
        ws.cell(row=cond_row, column=2, value="Tiempo de Entrega:").font = bold_font
        ws.cell(row=cond_row, column=3, value=cotizacion.tiempo_entrega or "A convenir")
        
        cond_row += 1
        ws.cell(row=cond_row, column=2, value="Forma de Pago:").font = bold_font
        ws.cell(row=cond_row, column=3, value=cotizacion.forma_pago or "Contado")
        
        cond_row += 1
        ws.cell(row=cond_row, column=2, value="Garantía:").font = bold_font
        ws.cell(row=cond_row, column=3, value=cotizacion.garantia or "1 Año")
        
        cond_row += 1
        ws.cell(row=cond_row, column=2, value="Validez de Oferta:").font = bold_font
        ws.cell(row=cond_row, column=3, value=cotizacion.validez_oferta or "30 Días")
        
        # Signature block
        sign_row = cond_row + 2
        ws.cell(row=sign_row, column=2, value="Atentamente,").font = bold_font
        
        firma_path = os.path.join(settings.BASE_DIR, 'firma_kave.png')
        if os.path.exists(firma_path):
            try:
                img_firma = XLImage(firma_path)
                img_firma.width = 180
                img_firma.height = 80
                ws.add_image(img_firma, f'B{sign_row + 1}')
            except Exception:
                pass
        
        sign_row += 6
        ws.cell(row=sign_row, column=2, value="Ing Alejandro Ochoa P.")
        sign_row += 1
        ws.cell(row=sign_row, column=2, value="Gerente")
        sign_row += 2
        ws.cell(row=sign_row, column=2, value="Calle 20 sur No 35-120")
        sign_row += 1
        ws.cell(row=sign_row, column=2, value="Medellin-Colombia")
        sign_row += 1
        ws.cell(row=sign_row, column=2, value="Tel: 300 7849624")
        sign_row += 1
        ws.cell(row=sign_row, column=2, value="aochoa@transformadoreskave.com")
        sign_row += 1
        ws.cell(row=sign_row, column=2, value="www.transformadoreskave.com")
        
        # Prepare response
        response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
        response['Content-Disposition'] = f'attachment; filename=Cotizacion_{cotizacion.numero_cotizacion or cotizacion.id}.xlsx'
        wb.save(response)
        
        return response
