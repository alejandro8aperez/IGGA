"""PDF & Excel exports replicating the original Informe Diario layout."""
import os
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side

from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib import colors
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import cm
from reportlab.platypus import (
    SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer, PageBreak, Image,
)


# --------------------------------------------------------------------------
# Excel export
# --------------------------------------------------------------------------

def _write_cell(ws, row, col, value, font=None, fill=None, alignment=None, border=None, number_format=None):
    cell = ws.cell(row, col, value)
    if font: cell.font = font
    if fill: cell.fill = fill
    if alignment: cell.alignment = alignment
    if border: cell.border = border
    if number_format: cell.number_format = number_format
    return cell


def _merge_and_write(ws, row, col_start, col_end, value, font=None, fill=None, alignment=None, border=None):
    """Merge cells and write value/styles to the anchor (top-left) cell only."""
    ws.merge_cells(start_row=row, start_column=col_start, end_row=row, end_column=col_end)
    return _write_cell(ws, row, col_start, value, font, fill, alignment, border)


def _fill_range(ws, row, col_start, col_end, font=None, fill=None, alignment=None, border=None):
    """Pre-fill a range of cells with styles BEFORE merging."""
    for c in range(col_start, col_end + 1):
        _write_cell(ws, row, c, '', font=font, fill=fill, alignment=alignment, border=border)


def generar_excel(informe) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = 'INFORME DIARIO'
    ws.sheet_properties.tabColor = '1E3A8A'

    NAVY = '1E3A8A'
    LIGHT_BLUE = 'DBEAFE'
    VERY_LIGHT = 'F8FAFC'
    WHITE = 'FFFFFF'
    BORDER_COLOR = 'CBD5E1'
    MUTED_TEXT = '64748B'

    title_font = Font(bold=True, size=14, color=NAVY, name='Calibri')
    subtitle_font = Font(size=9, color=MUTED_TEXT, name='Calibri')
    label_font = Font(bold=True, size=10, color='475569', name='Calibri')
    value_font = Font(bold=True, size=10, color='1E293B', name='Calibri')
    section_font = Font(bold=True, size=10, color=WHITE, name='Calibri')
    body_font = Font(size=10, color='1E293B', name='Calibri')
    body_bold = Font(bold=True, size=10, color='1E293B', name='Calibri')
    total_font = Font(bold=True, size=10, color=NAVY, name='Calibri')
    signature_label = Font(bold=True, size=9, color='475569', name='Calibri')
    signature_name = Font(bold=True, size=10, color='1E293B', name='Calibri')
    signature_cargo = Font(size=9, color=MUTED_TEXT, name='Calibri')

    navy_fill = PatternFill('solid', fgColor=NAVY)
    light_blue_fill = PatternFill('solid', fgColor=LIGHT_BLUE)
    light_fill = PatternFill('solid', fgColor=VERY_LIGHT)
    white_fill = PatternFill('solid', fgColor=WHITE)

    thin_side = Side(border_style='thin', color=BORDER_COLOR)
    cell_border = Border(left=thin_side, right=thin_side, top=thin_side, bottom=thin_side)
    no_border = Border()

    center_wrap = Alignment(horizontal='center', vertical='center', wrap_text=True)
    left_wrap = Alignment(horizontal='left', vertical='center', wrap_text=True)
    center_mid = Alignment(horizontal='center', vertical='center')

    col_widths = {
        'A': 4.5, 'B': 32, 'C': 12, 'D': 14, 'E': 14,
        'F': 4.5, 'G': 12, 'H': 14, 'I': 14, 'J': 4,
    }
    for letter, width in col_widths.items():
        ws.column_dimensions[letter].width = width

    row = 1

    # ═══════════════════════════════════════════════════════════════════════
    # HEADER
    # ═══════════════════════════════════════════════════════════════════════
    _merge_and_write(ws, row, 1, 9, 'LIBRO DIARIO DE OBRA — INTERVENTORÍA', title_font, alignment=center_wrap)
    row += 1

    cod = informe.codigo_formato or 'F-141-IN'
    _merge_and_write(ws, row, 1, 9,
                     f'Código: {cod}  |  Emisión: 27/08/2009  |  Mod: 00  |  Versión: 1',
                     subtitle_font, alignment=center_wrap)
    row += 2

    # ═══════════════════════════════════════════════════════════════════════
    # INFO TABLE
    # ═══════════════════════════════════════════════════════════════════════
    proyecto_nombre = informe.proyecto.nombre if informe.proyecto else '—'
    cliente_nombre = informe.proyecto.cliente.nombre if informe.proyecto and informe.proyecto.cliente else '—'
    fecha_str = informe.fecha.strftime('%d/%m/%Y') if informe.fecha else '—'
    dia_semana = informe.dia_semana or '—'
    status = informe.get_status_display() if hasattr(informe, 'get_status_display') else (informe.status or '—')
    topografia = 'Sí' if informe.comision_topografia else 'No'

    # Row 1: OBRA
    info_start = row
    _write_cell(ws, row, 1, 'OBRA:', label_font, light_fill, left_wrap, cell_border)
    _merge_and_write(ws, row, 2, 9, proyecto_nombre, value_font, white_fill, left_wrap, cell_border)
    row += 1

    # Row 2: CLIENTE | FECHA | ESTADO
    _write_cell(ws, row, 1, 'CLIENTE:', label_font, light_fill, left_wrap, cell_border)
    _merge_and_write(ws, row, 2, 3, cliente_nombre, value_font, white_fill, left_wrap, cell_border)
    _write_cell(ws, row, 4, 'FECHA:', label_font, light_fill, center_wrap, cell_border)
    _write_cell(ws, row, 5, fecha_str, value_font, white_fill, center_wrap, cell_border)
    _write_cell(ws, row, 6, 'ESTADO:', label_font, light_fill, center_wrap, cell_border)
    _merge_and_write(ws, row, 7, 9, status, value_font, white_fill, center_wrap, cell_border)
    row += 1

    # Row 3: DÍA | COM. TOPOGRAFÍA
    _write_cell(ws, row, 1, 'DÍA:', label_font, light_fill, left_wrap, cell_border)
    _merge_and_write(ws, row, 2, 3, dia_semana, value_font, white_fill, left_wrap, cell_border)
    _write_cell(ws, row, 4, 'COM. TOPOGRAFÍA:', label_font, light_fill, center_wrap, cell_border)
    _merge_and_write(ws, row, 5, 6, topografia, value_font, white_fill, center_wrap, cell_border)
    for c in range(7, 10):
        _write_cell(ws, row, c, '', border=cell_border)
    row += 2

    # ═══════════════════════════════════════════════════════════════════════
    # RAIN REPORT
    # ═══════════════════════════════════════════════════════════════════════
    _merge_and_write(ws, row, 1, 9, 'REPORTE DE LLUVIA',
                     Font(bold=True, size=10, color=NAVY, name='Calibri'), alignment=left_wrap)
    row += 1

    horas_lluvia = {r.hora: r.con_lluvia for r in informe.reportes_lluvia.all()}
    rain_count = sum(1 for h in range(24) if horas_lluvia.get(h))

    # Header row
    rain_font = Font(bold=True, size=8, color=WHITE, name='Calibri')
    _write_cell(ws, row, 1, 'Hora', rain_font, navy_fill, center_mid, cell_border)
    for h in range(24):
        _write_cell(ws, row, h + 2, h, rain_font, navy_fill, center_mid, cell_border)
    row += 1

    # Data row
    _write_cell(ws, row, 1, 'Lluvia', rain_font, navy_fill, center_mid, cell_border)
    for h in range(24):
        c = h + 2
        has_rain = horas_lluvia.get(h, False)
        fill = PatternFill('solid', fgColor='3B82F6') if has_rain else light_fill
        text_color = WHITE if has_rain else '94A3B8'
        _write_cell(ws, row, c, '///' if has_rain else '---',
                    Font(size=8, bold=True, color=text_color, name='Calibri'),
                    fill, center_mid, cell_border)
    row += 1

    _merge_and_write(ws, row, 1, 9, f'Horas con lluvia: {rain_count} de 24',
                     Font(size=9, color=MUTED_TEXT, name='Calibri'), alignment=left_wrap)
    row += 2

    # ═══════════════════════════════════════════════════════════════════════
    # TERRAIN STATUS
    # ═══════════════════════════════════════════════════════════════════════
    inicio = informe.estado_terreno_inicio or '—'
    final = informe.estado_terreno_final or '—'

    _write_cell(ws, row, 1, 'Estado Terreno Inicio:', body_bold, alignment=left_wrap)
    _merge_and_write(ws, row, 2, 4, inicio, body_font, alignment=left_wrap)
    _write_cell(ws, row, 5, '', border=no_border)
    _write_cell(ws, row, 6, 'Estado Terreno Final:', body_bold, alignment=left_wrap)
    _merge_and_write(ws, row, 7, 9, final, body_font, alignment=left_wrap)
    row += 2

    # ═══════════════════════════════════════════════════════════════════════
    # MAQUINARIA
    # ═══════════════════════════════════════════════════════════════════════
    detalles = list(informe.detalles.select_related('recurso__categoria'))
    maquinaria = [d for d in detalles if not d.recurso.categoria.nombre.upper().startswith('PERSONAL')]
    personal = [d for d in detalles if d.recurso.categoria.nombre.upper().startswith('PERSONAL')]
    maquinaria += list(informe.maquinaria_libre.all().order_by('orden'))
    personal += list(informe.personal_libre.all().order_by('orden'))

    _fill_range(ws, row, 1, 9, section_font, navy_fill, left_wrap, cell_border)
    _merge_and_write(ws, row, 1, 9, 'MAQUINARIA — EQUIPOS — HERRAMIENTAS DE PODER Y VEHÍCULOS',
                     section_font, navy_fill, left_wrap, cell_border)
    row += 1

    _write_cell(ws, row, 1, 'DESCRIPCIÓN', section_font, navy_fill, left_wrap, cell_border)
    _write_cell(ws, row, 2, 'CANT.', section_font, navy_fill, center_mid, cell_border)
    _write_cell(ws, row, 3, 'EMPRESA', section_font, navy_fill, center_mid, cell_border)
    _fill_range(ws, row, 4, 9, section_font, navy_fill, left_wrap, cell_border)
    _merge_and_write(ws, row, 4, 9, 'NOTAS', section_font, navy_fill, left_wrap, cell_border)
    row += 1

    total_maq = 0
    for i, item in enumerate(maquinaria):
        fill = light_fill if i % 2 == 1 else white_fill
        nombre = item.recurso.nombre if hasattr(item, 'recurso') else item.descripcion
        cant = float(item.cantidad) if item.cantidad else 0
        total_maq += cant
        empresa = getattr(item, 'empresa', '') or '—'
        notas = getattr(item, 'notas', '') or '—'

        _write_cell(ws, row, 1, nombre, body_font, fill, left_wrap, cell_border)
        _write_cell(ws, row, 2, cant, body_bold, fill, center_mid, cell_border)
        _write_cell(ws, row, 3, empresa, body_font, fill, left_wrap, cell_border)
        _fill_range(ws, row, 4, 9, Font(size=9, color=MUTED_TEXT, name='Calibri'), fill, left_wrap, cell_border)
        _merge_and_write(ws, row, 4, 9, notas, Font(size=9, color=MUTED_TEXT, name='Calibri'), fill, left_wrap, cell_border)
        row += 1

    if not maquinaria:
        _fill_range(ws, row, 1, 9, Font(size=9, color=MUTED_TEXT, italic=True, name='Calibri'), alignment=center_wrap, border=cell_border)
        _merge_and_write(ws, row, 1, 9, 'Sin maquinaria registrada',
                         Font(size=9, color=MUTED_TEXT, italic=True, name='Calibri'), alignment=center_wrap, border=cell_border)
        row += 1

    _fill_range(ws, row, 1, 3, border=cell_border)
    _merge_and_write(ws, row, 1, 3, 'TOTAL MAQUINARIA', total_font, light_blue_fill, left_wrap, cell_border)
    _fill_range(ws, row, 4, 9, border=cell_border)
    _merge_and_write(ws, row, 4, 9, total_maq, Font(bold=True, size=11, color=NAVY, name='Calibri'),
                     light_blue_fill, center_mid, cell_border)
    row += 2

    # ═══════════════════════════════════════════════════════════════════════
    # PERSONAL
    # ═══════════════════════════════════════════════════════════════════════
    _fill_range(ws, row, 1, 9, section_font, navy_fill, left_wrap, cell_border)
    _merge_and_write(ws, row, 1, 9, 'PERSONAL DE OBRA', section_font, navy_fill, left_wrap, cell_border)
    row += 1

    _write_cell(ws, row, 1, 'CARGO / DESCRIPCIÓN', section_font, navy_fill, left_wrap, cell_border)
    _write_cell(ws, row, 2, 'CANT.', section_font, navy_fill, center_mid, cell_border)
    _write_cell(ws, row, 3, 'EMPRESA', section_font, navy_fill, center_mid, cell_border)
    _fill_range(ws, row, 4, 9, section_font, navy_fill, left_wrap, cell_border)
    _merge_and_write(ws, row, 4, 9, 'NOTAS', section_font, navy_fill, left_wrap, cell_border)
    row += 1

    total_per = 0
    for i, item in enumerate(personal):
        fill = light_fill if i % 2 == 1 else white_fill
        nombre = item.recurso.nombre if hasattr(item, 'recurso') else item.descripcion
        cant = float(item.cantidad) if item.cantidad else 0
        total_per += cant
        empresa = getattr(item, 'empresa', '') or '—'
        notas = getattr(item, 'notas', '') or '—'

        _write_cell(ws, row, 1, nombre, body_font, fill, left_wrap, cell_border)
        _write_cell(ws, row, 2, cant, body_bold, fill, center_mid, cell_border)
        _write_cell(ws, row, 3, empresa, body_font, fill, left_wrap, cell_border)
        _fill_range(ws, row, 4, 9, Font(size=9, color=MUTED_TEXT, name='Calibri'), fill, left_wrap, cell_border)
        _merge_and_write(ws, row, 4, 9, notas, Font(size=9, color=MUTED_TEXT, name='Calibri'), fill, left_wrap, cell_border)
        row += 1

    if not personal:
        _fill_range(ws, row, 1, 9, Font(size=9, color=MUTED_TEXT, italic=True, name='Calibri'), alignment=center_wrap, border=cell_border)
        _merge_and_write(ws, row, 1, 9, 'Sin personal registrado',
                         Font(size=9, color=MUTED_TEXT, italic=True, name='Calibri'), alignment=center_wrap, border=cell_border)
        row += 1

    _fill_range(ws, row, 1, 3, border=cell_border)
    _merge_and_write(ws, row, 1, 3, 'TOTAL PERSONAL', total_font, light_blue_fill, left_wrap, cell_border)
    _fill_range(ws, row, 4, 9, border=cell_border)
    _merge_and_write(ws, row, 4, 9, total_per, Font(bold=True, size=11, color=NAVY, name='Calibri'),
                     light_blue_fill, center_mid, cell_border)
    row += 2

    # ═══════════════════════════════════════════════════════════════════════
    # ACTIVIDADES
    # ═══════════════════════════════════════════════════════════════════════
    _fill_range(ws, row, 1, 9, section_font, navy_fill, left_wrap, cell_border)
    _merge_and_write(ws, row, 1, 9, 'ACTIVIDADES DEL DÍA', section_font, navy_fill, left_wrap, cell_border)
    row += 1

    cats = {}
    for act in informe.actividades.select_related('categoria').order_by('categoria__orden', 'orden'):
        cats.setdefault(act.categoria.nombre, []).append(act.descripcion)

    if not cats:
        _fill_range(ws, row, 1, 9, Font(size=9, color=MUTED_TEXT, italic=True, name='Calibri'), alignment=center_wrap, border=cell_border)
        _merge_and_write(ws, row, 1, 9, 'Sin actividades registradas',
                         Font(size=9, color=MUTED_TEXT, italic=True, name='Calibri'), alignment=center_wrap, border=cell_border)
        row += 1
    else:
        for cat_nombre, items in cats.items():
            _write_cell(ws, row, 1, '', fill=PatternFill('solid', fgColor=NAVY), border=cell_border)
            _fill_range(ws, row, 2, 9, Font(bold=True, size=10, color='475569', name='Calibri'), light_fill, left_wrap, cell_border)
            _merge_and_write(ws, row, 2, 9, cat_nombre,
                             Font(bold=True, size=10, color='475569', name='Calibri'), light_fill, left_wrap, cell_border)
            row += 1

            for idx, item in enumerate(items, 1):
                _merge_and_write(ws, row, 1, 9, f'{idx}. {item}', body_font, alignment=left_wrap)
                row += 1
            row += 1

    row += 1

    # ═══════════════════════════════════════════════════════════════════════
    # OBSERVACIONES
    # ═══════════════════════════════════════════════════════════════════════
    _merge_and_write(ws, row, 1, 9, 'OBSERVACIONES GENERALES',
                     Font(bold=True, size=10, color=NAVY, name='Calibri'), alignment=left_wrap)
    row += 1

    obs = informe.observaciones_generales or 'Sin observaciones'
    obs_border = Border(left=Side(style='thin', color=NAVY),
                        right=Side(style='thin', color=BORDER_COLOR),
                        top=Side(style='thin', color=BORDER_COLOR),
                        bottom=Side(style='thin', color=BORDER_COLOR))
    obs_border_rest = Border(left=Side(style='thin', color=BORDER_COLOR),
                             right=Side(style='thin', color=BORDER_COLOR),
                             top=Side(style='thin', color=BORDER_COLOR),
                             bottom=Side(style='thin', color=BORDER_COLOR))
    _fill_range(ws, row, 1, 9, Font(size=10, color='1E293B', name='Calibri'),
                PatternFill('solid', fgColor='FAFBFC'), left_wrap, obs_border_rest)
    _merge_and_write(ws, row, 1, 9, obs, Font(size=10, color='1E293B', name='Calibri'),
                     PatternFill('solid', fgColor='FAFBFC'), left_wrap, obs_border)
    ws.row_dimensions[row].height = max(30, min(120, len(obs or '') // 2))
    row += 2

    # ═══════════════════════════════════════════════════════════════════════
    # ITEMS DE OBRA
    # ═══════════════════════════════════════════════════════════════════════
    items_obra = list(informe.items_obra.all().order_by('orden'))
    if items_obra:
        _fill_range(ws, row, 1, 9, section_font, navy_fill, left_wrap, cell_border)
        _merge_and_write(ws, row, 1, 9, 'ÍTEMS DE OBRA', section_font, navy_fill, left_wrap, cell_border)
        row += 1

        _write_cell(ws, row, 1, 'ITEM', section_font, navy_fill, center_mid, cell_border)
        _fill_range(ws, row, 2, 5, section_font, navy_fill, left_wrap, cell_border)
        _merge_and_write(ws, row, 2, 5, 'DESCRIPCIÓN', section_font, navy_fill, left_wrap, cell_border)
        _fill_range(ws, row, 6, 8, section_font, navy_fill, center_mid, cell_border)
        _merge_and_write(ws, row, 6, 8, 'EMPRESA', section_font, navy_fill, center_mid, cell_border)
        _write_cell(ws, row, 9, 'CANT.', section_font, navy_fill, center_mid, cell_border)
        row += 1

        for i, it in enumerate(items_obra):
            fill = light_fill if i % 2 == 1 else white_fill
            _write_cell(ws, row, 1, it.item or '—', body_bold, fill, center_mid, cell_border)
            _fill_range(ws, row, 2, 5, body_font, fill, left_wrap, cell_border)
            _merge_and_write(ws, row, 2, 5, it.descripcion or '—', body_font, fill, left_wrap, cell_border)
            _fill_range(ws, row, 6, 8, body_font, fill, left_wrap, cell_border)
            _merge_and_write(ws, row, 6, 8, it.empresa or '—', body_font, fill, left_wrap, cell_border)
            _write_cell(ws, row, 9, float(it.cantidad) if it.cantidad else 0, body_bold, fill, center_mid, cell_border)
            row += 1

        row += 1

    # ═══════════════════════════════════════════════════════════════════════
    # FIRMAS
    # ═══════════════════════════════════════════════════════════════════════
    row += 1
    for c in range(1, 10):
        _write_cell(ws, row, c, '', border=Border(top=Side(style='medium', color=NAVY)))
    row += 1

    elaborado_nombre = '____________________'
    elaborado_cargo = ''
    revisado_nombre = '____________________'
    revisado_cargo = ''

    if hasattr(informe, 'elaborado_por') and informe.elaborado_por:
        emp = informe.elaborado_por
        try:
            nombre = (f"{emp.primer_nombre} {emp.primer_apellido}").strip()
            elaborado_nombre = nombre or str(emp)
        except Exception:
            elaborado_nombre = str(emp) or '____________________'
        try:
            elaborado_cargo = str(emp.cargo) if hasattr(emp, 'cargo') and emp.cargo else ''
        except Exception:
            elaborado_cargo = ''

    if hasattr(informe, 'revisado_por') and informe.revisado_por:
        emp = informe.revisado_por
        try:
            nombre = (f"{emp.primer_nombre} {emp.primer_apellido}").strip()
            revisado_nombre = nombre or str(emp)
        except Exception:
            revisado_nombre = str(emp) or '____________________'
        try:
            revisado_cargo = str(emp.cargo) if hasattr(emp, 'cargo') and emp.cargo else ''
        except Exception:
            revisado_cargo = ''

    _fill_range(ws, row, 1, 3, fill=light_fill, border=cell_border)
    _merge_and_write(ws, row, 1, 3, 'ELABORADO POR', signature_label, light_fill, center_mid, cell_border)
    _fill_range(ws, row, 4, 6, fill=light_fill, border=cell_border)
    _merge_and_write(ws, row, 4, 6, 'REVISADO POR', signature_label, light_fill, center_mid, cell_border)
    _fill_range(ws, row, 7, 9, fill=light_fill, border=cell_border)
    _merge_and_write(ws, row, 7, 9, 'APROBADO POR', signature_label, light_fill, center_mid, cell_border)
    row += 1

    _fill_range(ws, row, 1, 3, border=cell_border)
    _merge_and_write(ws, row, 1, 3, elaborado_nombre, signature_name, white_fill, center_mid, cell_border)
    _fill_range(ws, row, 4, 6, border=cell_border)
    _merge_and_write(ws, row, 4, 6, revisado_nombre, signature_name, white_fill, center_mid, cell_border)
    _fill_range(ws, row, 7, 9, border=cell_border)
    _merge_and_write(ws, row, 7, 9, '____________________', signature_name, white_fill, center_mid, cell_border)
    row += 1

    _fill_range(ws, row, 1, 3, border=cell_border)
    _merge_and_write(ws, row, 1, 3, elaborado_cargo, signature_cargo, white_fill, center_mid, cell_border)
    _fill_range(ws, row, 4, 6, border=cell_border)
    _merge_and_write(ws, row, 4, 6, revisado_cargo, signature_cargo, white_fill, center_mid, cell_border)
    _fill_range(ws, row, 7, 9, border=cell_border)
    _merge_and_write(ws, row, 7, 9, '', signature_cargo, white_fill, center_mid, cell_border)

    # ═══════════════════════════════════════════════════════════════════════
    # PRINT SETTINGS
    # ═══════════════════════════════════════════════════════════════════════
    ws.print_title_rows = '1:5'
    ws.page_setup.orientation = 'landscape'
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


# --------------------------------------------------------------------------
# PDF export
# --------------------------------------------------------------------------

def generar_pdf(informe) -> bytes:
    buf = BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=landscape(A4),
                            leftMargin=1 * cm, rightMargin=1 * cm,
                            topMargin=1 * cm, bottomMargin=1 * cm)
    styles = getSampleStyleSheet()
    HEADER_BG = colors.HexColor('#1e3a8a')
    h1 = ParagraphStyle('h1', parent=styles['Heading1'], fontSize=12,
                        alignment=1, spaceAfter=4)
    h2 = ParagraphStyle('h2', parent=styles['Heading3'], fontSize=9,
                        backColor=HEADER_BG, textColor=colors.white, leading=11,
                        spaceBefore=4, spaceAfter=2)
    body = ParagraphStyle('body', parent=styles['BodyText'], fontSize=8,
                          leading=10)

    flow = []

    logo_path = os.path.join(os.path.dirname(__file__), 'static', 'informe_diario', 'logo.png')
    if os.path.isfile(logo_path):
        logo_img = Image(logo_path, width=3.5*cm, height=1.5*cm)
        header_table = Table([[logo_img, Paragraph('CONSTRUCCIÓN DE OBRA – LIBRO DIARIO DE OBRA / INTERVENTORÍA', h1)]],
                             colWidths=[4*cm, 20*cm])
        header_table.setStyle(TableStyle([
            ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
        ]))
        flow.append(header_table)
    else:
        flow.append(Paragraph('CONSTRUCCIÓN DE OBRA – LIBRO DIARIO DE OBRA / INTERVENTORÍA', h1))

    info = [
        ['OBRA:', (informe.proyecto.nombre if informe.proyecto else '—'), 'FECHA:', (informe.fecha.strftime('%d/%m/%Y') if informe.fecha else '—'),
         'DÍA:', informe.dia_semana or '—'],
    ]
    t = Table(info, colWidths=[2 * cm, 10 * cm, 2 * cm, 3 * cm, 1.5 * cm, 3 * cm])
    t.setStyle(TableStyle([
        ('FONT', (0, 0), (-1, -1), 'Helvetica', 8),
        ('FONT', (0, 0), (0, 0), 'Helvetica-Bold', 8),
        ('FONT', (2, 0), (2, 0), 'Helvetica-Bold', 8),
        ('FONT', (4, 0), (4, 0), 'Helvetica-Bold', 8),
        ('GRID', (0, 0), (-1, -1), 0.4, colors.grey),
    ]))
    flow.append(t)
    flow.append(Spacer(1, 6))

    flow.append(Paragraph('REPORTE DE LLUVIA (horas con lluvia)', h2))
    horas_lluvia = {r.hora: r.con_lluvia for r in informe.reportes_lluvia.all()}
    fila_h = ['Hora'] + [str(h) for h in range(24)]
    fila_v = ['Lluvia'] + ['X' if horas_lluvia.get(h) else '' for h in range(24)]
    t = Table([fila_h, fila_v], colWidths=[1.6 * cm] + [0.95 * cm] * 24)
    t.setStyle(TableStyle([
        ('FONT', (0, 0), (-1, -1), 'Helvetica', 7),
        ('FONT', (0, 0), (-1, 0), 'Helvetica-Bold', 7),
        ('BACKGROUND', (0, 0), (-1, 0), HEADER_BG),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
        ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
        ('GRID', (0, 0), (-1, -1), 0.3, colors.grey),
    ]))
    flow.append(t)
    flow.append(Spacer(1, 6))

    detalles = list(informe.detalles.select_related('recurso__categoria'))
    maquinaria = [d for d in detalles if not d.recurso.categoria.nombre.upper().startswith('PERSONAL')]
    personal = [d for d in detalles if d.recurso.categoria.nombre.upper().startswith('PERSONAL')]

    maquinaria += list(informe.maquinaria_libre.all().order_by('orden'))
    personal += list(informe.personal_libre.all().order_by('orden'))

    max_rows = max(len(maquinaria), len(personal), 1)
    data = [['MAQUINARIA - EQUIPOS - HERRAMIENTAS - VEHÍCULOS', 'CANT.',
             'PERSONAL DE OBRA', 'CANT.']]
    for i in range(max_rows):
        if i < len(maquinaria):
            item = maquinaria[i]
            l_name = item.recurso.nombre if hasattr(item, 'recurso') else item.descripcion
            l_qty = str(item.cantidad)
        else:
            l_name = l_qty = ''
        if i < len(personal):
            item = personal[i]
            r_name = item.recurso.nombre if hasattr(item, 'recurso') else item.descripcion
            r_qty = str(item.cantidad)
        else:
            r_name = r_qty = ''
        data.append([l_name, l_qty, r_name, r_qty])

    data.append(['TOTAL',
                 str(sum(float(getattr(d, 'cantidad', 0)) for d in maquinaria)),
                 'Total Personal',
                 str(sum(float(getattr(d, 'cantidad', 0)) for d in personal))])
    t = Table(data, colWidths=[10 * cm, 2 * cm, 10 * cm, 2 * cm], repeatRows=1)
    t.setStyle(TableStyle([
        ('FONT', (0, 0), (-1, -1), 'Helvetica', 7),
        ('FONT', (0, 0), (-1, 0), 'Helvetica-Bold', 7),
        ('FONT', (0, -1), (-1, -1), 'Helvetica-Bold', 7),
        ('BACKGROUND', (0, 0), (-1, 0), HEADER_BG),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
        ('BACKGROUND', (0, -1), (-1, -1), HEADER_BG),
        ('TEXTCOLOR', (0, -1), (-1, -1), colors.white),
        ('GRID', (0, 0), (-1, -1), 0.3, colors.grey),
        ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
    ]))
    flow.append(t)
    flow.append(Spacer(1, 6))

    flow.append(Paragraph('OBSERVACIONES GENERALES', h2))
    flow.append(Paragraph((informe.observaciones_generales or '—').replace('\n', '<br/>'), body))
    flow.append(Spacer(1, 4))

    t = Table([
        ['ESTADO DEL TERRENO AL INICIO', 'ESTADO DEL TERRENO AL FINAL'],
        [Paragraph((informe.estado_terreno_inicio or '—').replace('\n', '<br/>'), body),
         Paragraph((informe.estado_terreno_final or '—').replace('\n', '<br/>'), body)],
    ], colWidths=[12 * cm, 12 * cm])
    t.setStyle(TableStyle([
        ('FONT', (0, 0), (-1, 0), 'Helvetica-Bold', 8),
        ('BACKGROUND', (0, 0), (-1, 0), HEADER_BG),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
        ('GRID', (0, 0), (-1, -1), 0.3, colors.grey),
        ('VALIGN', (0, 0), (-1, -1), 'TOP'),
    ]))
    flow.append(t)
    flow.append(Spacer(1, 6))

    cats = {}
    for act in informe.actividades.select_related('categoria').order_by('categoria__orden', 'orden'):
        cats.setdefault(act.categoria.nombre, []).append(act.descripcion)
    for cat_nombre, items in cats.items():
        flow.append(Paragraph(cat_nombre, h2))
        for idx, item in enumerate(items, 1):
            flow.append(Paragraph(f"{idx}. {item}".replace('\n', '<br/>'), body))
        flow.append(Spacer(1, 4))

    items_obra = list(informe.items_obra.all().order_by('orden'))
    if items_obra:
        flow.append(Paragraph('ÍTEMS DE OBRA', h2))
        data_items = [['Item', 'Descripción', 'Empresa', 'Cantidad']]
        for it in items_obra:
            data_items.append([
                Paragraph(it.item, body),
                Paragraph(it.descripcion.replace('\n', '<br/>'), body),
                Paragraph(it.empresa, body),
                Paragraph(str(float(it.cantidad)), body)
            ])
        t_items = Table(data_items, colWidths=[2 * cm, 12 * cm, 6 * cm, 4 * cm])
        t_items.setStyle(TableStyle([
            ('FONT', (0, 0), (-1, -1), 'Helvetica', 7),
            ('FONT', (0, 0), (-1, 0), 'Helvetica-Bold', 7),
            ('BACKGROUND', (0, 0), (-1, 0), HEADER_BG),
            ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
            ('GRID', (0, 0), (-1, -1), 0.3, colors.grey),
            ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
        ]))
        flow.append(t_items)
        flow.append(Spacer(1, 6))

    anexos = list(informe.anexos.all())
    if anexos:
        flow.append(PageBreak())
        flow.append(Paragraph('ANEXOS FOTOGRÁFICOS', h1))
        for a in anexos:
            try:
                img_path = a.imagen.url if a.imagen else None
                if img_path and img_path.startswith('http'):
                    flow.append(Paragraph(f"<b>Sección:</b> {a.get_seccion_display()}", body))
                    flow.append(Paragraph(f"<b>Descripción:</b> {a.descripcion}", body))
                    flow.append(Paragraph(f'<link href="{img_path}">{img_path}</link>', body))
                else:
                    flow.append(Paragraph(f"<b>{a.get_seccion_display()}:</b> {a.descripcion}", body))
                flow.append(Spacer(1, 4))
            except Exception:
                continue

    doc.build(flow)
    return buf.getvalue()
